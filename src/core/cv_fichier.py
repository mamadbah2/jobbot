"""Reconnaissance du type réel d'un fichier et extraction de son texte.

Module PUR : aucun réseau, aucune base, aucun framework. C'est lui qui porte le
filtre le plus rentable du produit — tout ce qui peut être refusé l'est ici,
**avant qu'un seul token soit payé** à DeepSeek.

Le type est déduit des octets de tête et jamais de l'extension : un fichier
renommé ne prouve rien. Et un DOCX *étant* une archive ZIP, même les octets de
tête ne suffisent pas à l'affirmer — d'où la distinction entre `CvTypeRefuse`
(les octets de tête disent non) et `CvIllisible` (ils disaient oui, mais le
corps est inexploitable : PDF chiffré, archive tronquée, ZIP sans document).
"""

from __future__ import annotations

import io
from collections.abc import Callable
from dataclasses import dataclass
from typing import Any, Literal

from src.core.erreurs import CvIllisible, CvSansTexte, CvTypeRefuse

TypeCv = Literal["pdf", "docx", "jpeg", "png"]

_SIGNATURES: tuple[tuple[bytes, TypeCv], ...] = (
    (b"%PDF", "pdf"),
    (b"\xff\xd8\xff", "jpeg"),
    (b"\x89PNG\r\n\x1a\n", "png"),
    # En dernier : une archive ZIP *peut* être un DOCX, mais aussi n'importe
    # quoi d'autre. C'est `extraire_texte` qui tranche.
    (b"PK\x03\x04", "docx"),
)

# Écart minimal entre deux caractères, en fraction du corps, au-delà duquel
# pdfplumber insère une espace. 0,15 sépare les mots LaTeX (~0,25 em) sans
# couper les crénages de Word (quelques millièmes d'em). Testé dans les deux sens.
_TOLERANCE_ESPACE = 0.15


def reconnaitre_type(contenu: bytes) -> TypeCv:
    for signature, type_cv in _SIGNATURES:
        if contenu.startswith(signature):
            return type_cv
    raise CvTypeRefuse


def est_une_image(type_cv: TypeCv) -> bool:
    """Une image n'a pas de texte à extraire localement : elle part telle quelle
    au modèle multimodal (spec §8)."""
    return type_cv in ("jpeg", "png")


def _lire_pdf(contenu: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(contenu)) as document:
        # Tolérance proportionnelle au corps du caractère, et non les 3 pt fixes
        # par défaut : LaTeX sépare les mots par un DÉCALAGE (~2,5 pt en corps
        # 10), sans caractère espace, et le défaut les fusionnait.
        return "\n".join(
            page.extract_text(x_tolerance_ratio=_TOLERANCE_ESPACE) or ""
            for page in document.pages
        )


def _lire_docx(contenu: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(contenu))
    morceaux = _textes_conteneur(document)
    # En-têtes et pieds de page : nom et coordonnées s'y trouvent souvent.
    for section in document.sections:
        morceaux += _textes_conteneur(section.header)
        morceaux += _textes_conteneur(section.footer)
    return "\n".join(morceaux)


def _textes_conteneur(conteneur: Any) -> list[str]:
    """Paragraphes puis tableaux d'un conteneur (document, en-tête, cellule).

    `document.paragraphs` ignore les tableaux, or la mise en page d'un CV en
    colonnes passe couramment par un tableau : le lire à moitié ferait refuser
    à tort un fichier plein de texte. Une cellule peut contenir un tableau,
    d'où la récursion. Les zones de texte flottantes ne sont pas lues.
    """
    textes = [p.text for p in conteneur.paragraphs]
    for tableau in conteneur.tables:
        for ligne in tableau.rows:
            for cellule in ligne.cells:
                textes += _textes_conteneur(cellule)
    return textes


@dataclass(frozen=True, slots=True)
class Lecteurs:
    """Lecteurs injectables.

    Les chemins d'échec (PDF chiffré, archive tronquée) se testent en injectant
    un lecteur qui lève, plutôt qu'en fabriquant et maintenant des fichiers
    corrompus dans le dépôt.
    """

    pdf: Callable[[bytes], str]
    docx: Callable[[bytes], str]


LECTEURS_REELS = Lecteurs(pdf=_lire_pdf, docx=_lire_docx)


def extraire_texte(
    contenu: bytes,
    type_cv: TypeCv,
    *,
    minimum: int,
    lecteurs: Lecteurs | None = None,
) -> str:
    """Rend le texte du document, ou lève.

    `minimum` compte les caractères **hors espaces** : un PDF scanné rend
    rarement zéro caractère, il rend quelques artefacts de police. « Vide » ne
    serait donc pas un test suffisant.
    """
    if est_une_image(type_cv):  # pragma: no cover — l'appelant ne doit pas le faire
        raise CvIllisible

    lecteurs = lecteurs or LECTEURS_REELS
    lecteur = lecteurs.pdf if type_cv == "pdf" else lecteurs.docx

    lecture_echouee = False
    texte = ""
    try:
        texte = lecteur(contenu)
    except (CvSansTexte, CvIllisible):
        raise
    except Exception:
        # Une bibliothèque de lecture lève n'importe quoi sur un fichier
        # malformé (BadZipFile, KeyError, ValueError, OSError, erreurs de
        # pdfminer...) : tout échec de lecture est un refus, pas une panne.
        # On ne fait que lever un drapeau ICI : `raise` dans le bloc laisserait
        # l'exception d'origine, et son message, dans `__context__` même avec
        # `from None`. Le message peut contenir un fragment du CV, et une
        # alerte admin recopie tout son contexte par le réseau (§7).
        lecture_echouee = True
    if lecture_echouee:
        raise CvIllisible

    # `split()` coupe sur tous les blancs (sauts de ligne, tabulations), pas
    # seulement l'espace : un scan rend surtout des retours à la ligne.
    if len("".join(texte.split())) < minimum:
        raise CvSansTexte
    return texte
