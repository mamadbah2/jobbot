"""Reconnaissance de type et extraction de texte.

Les fixtures sont FABRIQUÉES ici, pas committées : un PDF minimal tient en une
chaîne d'octets, et `python-docx` sait écrire un DOCX. Rien de binaire à
maintenir dans le dépôt, et les cas limites sont reproductibles.
"""

from __future__ import annotations

import io
import traceback
import zipfile

import pytest

from src.core.cv_fichier import (
    Lecteurs,
    est_une_image,
    extraire_texte,
    reconnaitre_type,
)
from src.core.erreurs import CvIllisible, CvSansTexte, CvTypeRefuse

JPEG = b"\xff\xd8\xff\xe0" + b"\x00" * 64
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


def _docx(texte: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(texte)
    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()


def _zip_qui_nest_pas_un_docx() -> bytes:
    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w") as z:
        z.writestr("quelconque.txt", "ceci est une archive, pas un document Word")
    return tampon.getvalue()


def test_reconnait_les_quatre_types_acceptes() -> None:
    assert reconnaitre_type(b"%PDF-1.7\n...") == "pdf"
    assert reconnaitre_type(_docx("bonjour")) == "docx"
    assert reconnaitre_type(JPEG) == "jpeg"
    assert reconnaitre_type(PNG) == "png"


def test_refuse_ce_qui_nest_pas_un_type_accepte() -> None:
    """L'extension ne prouve rien : on ne regarde que les octets de tête."""
    with pytest.raises(CvTypeRefuse):
        reconnaitre_type(b"MZ\x90\x00 un executable renomme en .pdf")
    with pytest.raises(CvTypeRefuse):
        reconnaitre_type(b"")


def test_les_images_sont_signalees_comme_telles() -> None:
    assert est_une_image("jpeg") is True
    assert est_une_image("png") is True
    assert est_une_image("pdf") is False
    assert est_une_image("docx") is False


def test_extrait_le_texte_dun_docx() -> None:
    contenu = _docx("Awa Diop — Comptable — 8 ans d'expérience " * 10)
    texte = extraire_texte(contenu, "docx", minimum=50)
    assert "Awa Diop" in texte


def test_refuse_un_document_trop_pauvre_en_texte() -> None:
    """Le cas du PDF scanné : quelques artefacts de police, pas un CV.
    Refusé AVANT tout appel payant."""
    with pytest.raises(CvSansTexte):
        extraire_texte(_docx("abc"), "docx", minimum=200)


def test_un_zip_qui_nest_pas_un_docx_est_refuse_proprement() -> None:
    """Review Focus n°2 : un DOCX *est* un ZIP, donc les octets de tête ne
    prouvent rien. L'ouverture doit lever CvIllisible, pas remonter une
    exception de bibliothèque jusqu'au 500."""
    contenu = _zip_qui_nest_pas_un_docx()
    assert reconnaitre_type(contenu) == "docx"  # les octets de tête y croient
    with pytest.raises(CvIllisible):
        extraire_texte(contenu, "docx", minimum=10)


def test_un_pdf_chiffre_ou_tronque_est_refuse_proprement() -> None:
    """Review Focus n°1 : pdfplumber lève sur un PDF chiffré ou tronqué.
    L'utilisateur doit lire un refus, pas récolter un 500.

    On injecte un lecteur qui lève, au lieu de fabriquer un PDF chiffré :
    ce qu'on teste est la traduction de l'échec, pas pdfplumber.
    """

    def lecteur_qui_leve(_: bytes) -> str:
        raise ValueError("PDF starts with ... (encrypted)")

    with pytest.raises(CvIllisible):
        extraire_texte(
            b"%PDF-1.7 corps casse",
            "pdf",
            minimum=10,
            lecteurs=Lecteurs(pdf=lecteur_qui_leve, docx=lambda _: ""),
        )


def test_le_message_dexception_ne_porte_jamais_de_contenu_de_cv() -> None:
    """Règle n°1 de la phase : un fragment de CV ne doit pas pouvoir remonter
    dans une exception, qui finirait dans une alerte admin par le réseau (§7)."""
    secret = "Awa Diop née le 3 mai 1994 à Thiès"

    def lecteur_qui_fuit(_: bytes) -> str:
        raise ValueError(f"échec en lisant : {secret}")

    with pytest.raises(CvIllisible) as capture:
        extraire_texte(
            b"%PDF-1.7",
            "pdf",
            minimum=10,
            lecteurs=Lecteurs(pdf=lecteur_qui_fuit, docx=lambda _: ""),
        )
    assert secret not in str(capture.value)
    assert secret not in repr(capture.value)
    # `from None` masque l'affichage mais laisse l'exception d'origine dans
    # `__context__` si on lève dans le bloc `except` : on vérifie la chaîne.
    assert capture.value.__context__ is None
    assert capture.value.__cause__ is None
    assert secret not in "".join(traceback.format_exception(capture.value))


def test_le_seuil_ignore_tous_les_blancs() -> None:
    """Cas discriminant : longueur brute au-dessus du minimum, compte utile en
    dessous. Un scan rend surtout des sauts de ligne et des tabulations."""
    texte = "a" + "\n" * 10 + "\t" * 5 + " " * 5 + "b"
    assert len(texte) > 8
    with pytest.raises(CvSansTexte):
        extraire_texte(
            b"%PDF",
            "pdf",
            minimum=8,
            lecteurs=Lecteurs(pdf=lambda _: texte, docx=lambda _: ""),
        )


def test_borne_exacte_du_seuil() -> None:
    def avec(texte: str) -> Lecteurs:
        return Lecteurs(pdf=lambda _: texte, docx=lambda _: "")

    assert extraire_texte(b"%PDF", "pdf", minimum=5, lecteurs=avec("ab cd\ne")) == "ab cd\ne"
    with pytest.raises(CvSansTexte):
        extraire_texte(b"%PDF", "pdf", minimum=5, lecteurs=avec("ab c\nd"))


def test_un_docx_mis_en_page_en_tableau_est_lu() -> None:
    """Colonnes par tableau : courant dans les CV. Tableaux imbriqués inclus."""
    from docx import Document

    doc = Document()
    tableau = doc.add_table(rows=1, cols=2)
    tableau.cell(0, 0).text = "Awa Diop comptable huit ans d'experience"
    interne = tableau.cell(0, 1).add_table(rows=1, cols=1)
    interne.cell(0, 0).text = "Competences Sage Excel fiscalite"
    tampon = io.BytesIO()
    doc.save(tampon)
    texte = extraire_texte(tampon.getvalue(), "docx", minimum=50)
    assert "Awa Diop" in texte
    assert "Sage Excel" in texte


def _pdf(flux: bytes) -> bytes:
    """PDF minimal, police Helvetica standard (aucune police embarquée)."""
    objets = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842]"
        b" /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(flux) + flux + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    sortie = b"%PDF-1.4\n"
    positions = []
    for numero, objet in enumerate(objets, start=1):
        positions.append(len(sortie))
        sortie += b"%d 0 obj\n" % numero + objet + b"\nendobj\n"
    debut_xref = len(sortie)
    sortie += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objets) + 1)
    sortie += b"".join(b"%010d 00000 n \n" % p for p in positions)
    sortie += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objets) + 1,
        debut_xref,
    )
    return sortie


def test_un_pdf_latex_garde_ses_espaces_entre_les_mots() -> None:
    """LaTeX (xdvipdfmx) n'écrit aucun caractère espace : il DÉCALE le mot suivant.

    Avec la tolérance par défaut de pdfplumber (3 pt), un écart de 2,5 pt en
    corps 10 était fusionné : « ResponsableTechniqueIT ». Constaté sur 16 CV
    réels le 2026-10-21 (496 mots collés), que DeepSeek ne recollait pas
    toujours (« fleetprovisioning » rendu comme une compétence).
    """
    flux = b"BT /F1 10 Tf 72 700 Td [(Responsable) -250 (Technique) -250 (IT)] TJ ET"
    texte = extraire_texte(_pdf(flux), "pdf", minimum=5)
    assert "Responsable Technique IT" in texte


def test_le_crenage_dun_pdf_word_ne_coupe_pas_les_mots() -> None:
    """Garde-fou inverse : un crénage serré à l'intérieur d'un mot reste un mot.

    Word et LibreOffice ajustent l'approche entre lettres (quelques millièmes
    de cadratin) ; une tolérance trop serrée rendrait « Compt abilité ».
    """
    flux = b"BT /F1 10 Tf 72 700 Td [(Compt) 30 (abilit) -20 (\xe9 g\xe9n\xe9rale)] TJ ET"
    texte = extraire_texte(_pdf(flux), "pdf", minimum=5)
    assert "Comptabilité générale" in texte
